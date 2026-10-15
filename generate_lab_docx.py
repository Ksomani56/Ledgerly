"""
generate_lab_docx.py - Styling primitives, page setup, and structural helpers
for the Software Engineering & Web Development (WDL) Lab Manual.
Tailored to 'Ledgerly' Invoice & Quotation Generator PWA.
"""

import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

WORKSPACE = Path(r"C:\Users\Krushna\Documents\Projects\Invoice Generator")
OUTPUT_DIR = WORKSPACE / "Lab Manual Docs"
DIAGRAMS_DIR = WORKSPACE / "Project Diagrams"

OUTPUT_DIR.mkdir(exist_ok=True)

# Cohesive Palette (Academic & Corporate)
PRIMARY_COLOR = RGBColor(0x1B, 0x4D, 0x4B)   # Deep Teal #1B4D4B
SECONDARY_COLOR = RGBColor(0x2C, 0x3E, 0x50) # Slate #2C3E50
MUTED_COLOR = RGBColor(0x71, 0x80, 0x96)     # Gray #718096
DARK_INK = RGBColor(0x1A, 0x20, 0x2C)        # Body text #1A202C

HEADER_BG = "1B4D4B"                          # Table header fill
ROW_ALT_BG = "F7FAFC"                         # Alternating row fill
BORDER_COLOR = "CBD5E0"                       # Light table border
CALLOUT_BG = "F8FAFC"                         # Shaded note fill

def setup_document_page(doc):
    """Configures 1-inch margins and running academic header & footer."""
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        # Running Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Department of Computer Engineering  |  WDL Lab (CMPN Div C)")
        hrun.font.name = 'Calibri'
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = MUTED_COLOR

        # Running Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        frun = fp.add_run("Project Case Study: Ledgerly PWA\t\tAcademic Year 2026-27")
        frun.font.name = 'Calibri'
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = MUTED_COLOR

def set_cell_background(cell, hex_color):
    """Sets background fill color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding (in dxa) for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="000000", sz="8", val="single"):
    """Sets solid borders around a table and all its interior gridlines."""
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for b in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{b}')
        border.set(qn('w:val'), val)
        border.set(qn('w:sz'), sz)
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tblBorders.append(border)
    tblPr.append(tblBorders)

def style_heading(doc, text, level=1):
    """Adds a professionally formatted heading with exact color hierarchy."""
    h = doc.add_heading(text, level=level)
    run = h.runs[0]
    run.font.name = 'Calibri'
    if level == 1:
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = PRIMARY_COLOR
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
    elif level == 2:
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = SECONDARY_COLOR
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(3)
    elif level == 3:
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x34, 0x49, 0x5E)
        h.paragraph_format.space_before = Pt(6)
        h.paragraph_format.space_after = Pt(2)
    return h

def add_p(doc, text, bold_prefix=None, space_after=4.5):
    """Adds a standard body paragraph with clean academic typography."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = DARK_INK
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10.5)
    r.font.color.rgb = DARK_INK
    return p

def add_bullet(doc, text, bold_prefix=None):
    """Adds a clean bullet point."""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = DARK_INK
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10.5)
    r.font.color.rgb = DARK_INK
    return p

def add_student_info_table(doc):
    """Adds the required Student & Faculty identification table at top of document."""
    info_data = [
        ("NAME", "Kimaya Teli\nAbhishek kumar Kairamkonda\nKrishna Somani"),
        ("ROLL NO", "24102C0079\n24102C0085\n24102C0067"),
        ("BRANCH", "CMPN"),
        ("DIV", "C"),
        ("SUBJECT", "WDL LAB"),
        ("FACULTY", "Dr. Ravindra Sangle")
    ]
    table = doc.add_table(rows=len(info_data), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="000000", sz="8")

    col_widths = [Inches(1.8), Inches(4.7)]
    for row_idx, (label, val) in enumerate(info_data):
        row = table.rows[row_idx]

        # Label Cell
        c0 = row.cells[0]
        c0.width = col_widths[0]
        set_cell_margins(c0, top=100, bottom=100, left=140, right=140)
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        p0.paragraph_format.line_spacing = 1.15
        r0 = p0.add_run(label)
        r0.font.name = 'Calibri'
        r0.font.size = Pt(10.5)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        # Value Cell
        c1 = row.cells[1]
        c1.width = col_widths[1]
        set_cell_margins(c1, top=100, bottom=100, left=140, right=140)
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(val)
        r1.font.name = 'Calibri'
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    doc.add_paragraph() # spacer after info table

def add_header(doc, exp_num, title, module_num=None):
    """Initializes page margins, header, footer, identification table, and experiment title."""
    setup_document_page(doc)
    add_student_info_table(doc)

    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_top = p_top.add_run("SOFTWARE ENGINEERING & WEB DEVELOPMENT LAB MANUAL")
    r_top.font.name = 'Calibri'
    r_top.font.size = Pt(13)
    r_top.font.bold = True
    r_top.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    p_proj = doc.add_paragraph()
    p_proj.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_proj = p_proj.add_run("PROJECT CASE STUDY: LEDGERLY — INVOICE & QUOTATION GENERATOR PWA")
    r_proj.font.name = 'Calibri'
    r_proj.font.size = Pt(10)
    r_proj.font.italic = True
    r_proj.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    h = style_heading(doc, f"Experiment No. {exp_num}: {title}", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph() # spacer

def add_callout(doc, text, title="ENGINEERING STANDARD / COMPLIANCE NOTE"):
    """Adds a styled callout box with a thick accent left border."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')

    # Thick left border
    left_b = OxmlElement('w:left')
    left_b.set(qn('w:val'), 'single')
    left_b.set(qn('w:sz'), '24')
    left_b.set(qn('w:space'), '0')
    left_b.set(qn('w:color'), '1B4D4B')
    tblBorders.append(left_b)

    for b in ['top', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{b}')
        border.set(qn('w:val'), 'none')
        tblBorders.append(border)
    tblPr.append(tblBorders)

    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, CALLOUT_BG)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"{title}\n")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(10)
    r_title.font.bold = True
    r_title.font.color.rgb = PRIMARY_COLOR

    r_text = p.add_run(text)
    r_text.font.name = 'Calibri'
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = DARK_INK
    r_text.font.italic = True

    doc.add_paragraph() # spacer

def add_hardware_software_requirements(doc, hw_items=None, sw_items=None):
    """Renders a standard academic Hardware & Software specifications table."""
    style_heading(doc, "Hardware & Software Requirements", 2)
    default_hw = [
        ("Processor / CPU", "Intel Core i5 / AMD Ryzen 5 or higher (minimum 2.0 GHz)"),
        ("System Memory (RAM)", "8.0 GB RAM or higher recommended"),
        ("Storage / Disk", "256 GB SSD (Minimum 500 MB free workspace space)"),
        ("Display / Resolution", "1920 x 1080 Full HD (Responsive down to 360 x 640 mobile viewports)")
    ]
    default_sw = [
        ("Operating System", "Windows 11 / 10 (64-bit) / Ubuntu Linux / macOS"),
        ("Web Browser", "Google Chrome (v110+), Mozilla Firefox (v115+), Microsoft Edge"),
        ("Core Technologies", "HTML5, Vanilla CSS3 (Custom Variables), ECMAScript 2022 (Pure Vanilla JS)"),
        ("Runtime & Tooling", "Node.js LTS (v20+) for local asset bundling & Python 3.10+ for automated testing"),
        ("Testing Suite", "Python unittest & pytest automated testing harnesses"),
        ("Version Control", "Git v2.40+ & GitHub Pages CI/CD deployment")
    ]
    hw = hw_items or default_hw
    sw = sw_items or default_sw

    table_data = []
    for item in hw:
        table_data.append(["Hardware", item[0], item[1]])
    for item in sw:
        table_data.append(["Software / Runtime", item[0], item[1]])

    add_styled_table(doc, ["Category", "Component / Specification", "Minimum Recommended Configuration"], table_data,
                     col_widths=[Inches(1.5), Inches(2.2), Inches(2.8)])

def add_styled_table(doc, headers, data, col_widths=None):
    """Creates a clean, styled table with headers and alternate row background."""
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Style Header
    hdr_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        cell = hdr_cells[i]
        if col_widths and i < len(col_widths):
            cell.width = col_widths[i]
        set_cell_background(cell, HEADER_BG)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header_text)
        run.font.name = 'Calibri'
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Style Rows
    for row_idx, row_data in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        bg_col = ROW_ALT_BG if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            cell = row_cells[col_idx]
            if col_widths and col_idx < len(col_widths):
                cell.width = col_widths[col_idx]
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(str(cell_value))
            run.font.name = 'Calibri'
            run.font.size = Pt(9.5)
            run.font.color.rgb = DARK_INK

    # Light table border
    set_table_borders(table, color="D1D5DB", sz="4")
    doc.add_paragraph()
    return table

def add_code_block(doc, code_str, caption=None):
    """Adds a shaded monospace code box."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    set_table_borders(table, color="E2E8F0", sz="4")

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(code_str.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(9.0)
    run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    if caption:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(3)
        p_cap.paragraph_format.space_after = Pt(6)
        r_cap = p_cap.add_run(f"Listing: {caption}")
        r_cap.font.name = 'Calibri'
        r_cap.font.size = Pt(8.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = MUTED_COLOR
    else:
        doc.add_paragraph()

def add_image_if_exists(doc, img_name, caption=None, width=Inches(6.0)):
    """Inserts a diagram image with caption."""
    img_path = DIAGRAMS_DIR / img_name
    if img_path.exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(2)
        doc.add_picture(str(img_path), width=width)
        if caption:
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_cap = p_cap.add_run(f"Figure: {caption}")
            r_cap.font.name = 'Calibri'
            r_cap.font.size = Pt(9)
            r_cap.font.italic = True
            r_cap.font.color.rgb = MUTED_COLOR
        doc.add_paragraph()

def add_test_cases_table(doc, test_cases):
    """Renders a standard formal Test Cases & Verification table."""
    style_heading(doc, "Experimental Validation & Test Cases", 2)
    headers = ["Test ID", "Scenario / Feature Tested", "Input Data", "Expected Result", "Status"]
    add_styled_table(doc, headers, test_cases, col_widths=[Inches(1.0), Inches(1.8), Inches(1.5), Inches(1.5), Inches(0.7)])

def add_viva_questions(doc, qna_list):
    """Renders a rich Viva-Voce / Technical Interview review section."""
    style_heading(doc, "Review & Viva-Voce Questions with Solutions", 2)
    add_p(doc, "The following technical review questions assess core understanding of software engineering principles, web runtime behavior, and architecture applied in this experiment:")
    for idx, (q, a) in enumerate(qna_list, 1):
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(6)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        rq = p_q.add_run(f"Q{idx}. {q}")
        rq.font.name = 'Calibri'
        rq.font.size = Pt(10.5)
        rq.font.bold = True
        rq.font.color.rgb = PRIMARY_COLOR

        p_a = doc.add_paragraph()
        p_a.paragraph_format.space_before = Pt(0)
        p_a.paragraph_format.space_after = Pt(6)
        p_a.paragraph_format.line_spacing = 1.15
        ra_label = p_a.add_run("Answer: ")
        ra_label.font.name = 'Calibri'
        ra_label.font.size = Pt(10)
        ra_label.font.bold = True
        ra_label.font.color.rgb = SECONDARY_COLOR
        ra = p_a.add_run(a)
        ra.font.name = 'Calibri'
        ra.font.size = Pt(10)
        ra.font.color.rgb = DARK_INK

def add_conclusion(doc, text):
    """Renders the final conclusion and learning outcomes section."""
    style_heading(doc, "Conclusion & Learning Outcomes", 2)
    add_p(doc, text)
